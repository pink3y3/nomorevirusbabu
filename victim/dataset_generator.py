"""Generate a deterministic, entirely fictional military-themed lab dataset."""
from __future__ import annotations
import hashlib, json, os, random, struct, wave
from pathlib import Path
from datetime import datetime, timezone

try:
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image, ImageDraw
except ImportError as exc:
    raise SystemExit("Install dependencies first: pip install python-docx openpyxl pillow") from exc

try:
    import config
except ImportError:
    config = None

SEED = 20261009
TARGETS = {
    "01_Operations": {
        "Mission_Report_01.docx": "docx", "Deployment_Plan.xlsx": "xlsx",
        "Operations_Map.jpg": "jpg",
    },
    "02_Intelligence": {
        "Intelligence_Report.docx": "docx", "Personnel_Log.xlsx": "xlsx",
        "Surveillance_Image.jpg": "jpg",
    },
    "03_Media": {"Briefing_Audio.mp3": "mp3", "Training_Clip.mp4": "mp4"},
}
DECOYS = {
    "Mission_Orders_2026.docx": "docx", "Classified_Personnel.xlsx": "xlsx",
    "Strategic_Assessment.docx": "docx", "Satellite_Image.jpg": "jpg",
    "Briefing_Audio.mp3": "mp3", "Training_Clip.mp4": "mp4",
}

def _write_docx(path: Path, title: str, seed: int) -> None:
    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph("FICTIONAL TRAINING MATERIAL — NOT REAL MILITARY INFORMATION")
    doc.add_heading("Situation summary", level=1)
    doc.add_paragraph(f"Exercise reference: EX-{seed}; status: simulated; distribution: lab only.")
    doc.add_heading("Operational notes", level=1)
    for i in range(1, 7):
        doc.add_paragraph(
            f"Training note {i}: Unit {((seed + i) % 17) + 1:02d} reports a simulated "
            f"logistics checkpoint. No real locations, personnel, or operational details are used."
        )
    doc.add_heading("Review table", level=1)
    table = doc.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ["Item", "Status", "Exercise value"]):
        cell.text = value
    for i in range(1, 5):
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = f"Task {i}", "Simulated", str(seed + i)
    doc.save(path)

def _write_xlsx(path: Path, seed: int) -> None:
    rng = random.Random(seed)
    wb = Workbook()
    ws = wb.active
    ws.title = "Fictional Logistics"
    ws.append(["Exercise ID", "Fictional Unit", "Supply Type", "Quantity", "Status"])
    for i in range(1, 26):
        ws.append([f"EX-{seed}-{i:03d}", f"Unit-{rng.randint(1, 30):02d}",
                   rng.choice(["Food", "Medical", "Fuel", "Equipment"]),
                   rng.randint(10, 900), rng.choice(["Planned", "In transit", "Received"])])
    wb.save(path)

def _write_jpg(path: Path, seed: int) -> None:
    img = Image.new("RGB", (1024, 768), (35, 48, 65))
    draw = ImageDraw.Draw(img)
    draw.rectangle((45, 45, 979, 723), outline=(220, 190, 90), width=5)
    draw.text((80, 90), "FICTIONAL TRAINING IMAGE", fill="white")
    draw.text((80, 145), f"EXERCISE EX-{seed}", fill=(240, 210, 100))
    for i in range(8):
        x, y = 90 + (i * 103) % 800, 250 + (i * 57) % 380
        draw.rectangle((x, y, x + 45, y + 32), outline=(100, 200, 180), width=3)
    img.save(path, "JPEG", quality=88)

def _write_wav(path: Path, seed: int) -> None:
    # A short generated tone, not speech or real audio.
    rate, duration = 22050, 3
    freq = 330 + (seed % 220)
    with wave.open(str(path), "wb") as out:
        out.setnchannels(1); out.setsampwidth(2); out.setframerate(rate)
        frames = bytearray()
        for n in range(rate * duration):
            value = int(5000 * __import__("math").sin(2 * __import__("math").pi * freq * n / rate))
            frames.extend(struct.pack("<h", value))
        out.writeframes(frames)

def _write_mp3(path: Path, seed: int) -> None:
    import shutil, subprocess, tempfile
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        raise RuntimeError("ffmpeg is required to generate the .mp3 training clip. Install ffmpeg and rerun.")
    with tempfile.NamedTemporaryFile(suffix=".wav", delete=False) as tmp:
        wav_path = Path(tmp.name)
    try:
        _write_wav(wav_path, seed)
        subprocess.run([ffmpeg, "-y", "-loglevel", "error", "-i", str(wav_path),
                        "-codec:a", "libmp3lame", "-q:a", "7", str(path)], check=True)
    finally:
        wav_path.unlink(missing_ok=True)

def _write_mp4(path: Path, seed: int) -> None:
    # A small valid MP4 requires ffmpeg. Fail clearly instead of writing a fake .mp4.
    import shutil, subprocess
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        raise RuntimeError("ffmpeg is required to generate the .mp4 training clip. Install ffmpeg and rerun.")
    subprocess.run([ffmpeg, "-y", "-loglevel", "error", "-f", "lavfi", "-i",
                    "color=c=navy:s=640x360:d=2", "-vf",
                    f"drawtext=text='FICTIONAL EXERCISE {seed}':fontcolor=white:fontsize=24:x=30:y=150",
                    "-an", "-pix_fmt", "yuv420p", str(path)], check=True)

WRITERS = {"docx": _write_docx, "xlsx": _write_xlsx, "jpg": _write_jpg,
           "wav": _write_wav, "mp3": _write_mp3, "mp4": _write_mp4}

def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def generate_dataset(base_dir: str | Path) -> dict:
    """Create MILITARY_DATA and results/baseline_hashes.json beneath base_dir."""
    root = Path(base_dir).expanduser().resolve()
    data_dir = root if root.name == "MILITARY_DATA" else root / "MILITARY_DATA"
    # Safety: refuse to replace an existing dataset automatically.
    if data_dir.exists() and any(data_dir.iterdir()):
        raise FileExistsError(f"{data_dir} already contains files. Move/delete it manually before regenerating.")
    data_dir.mkdir(parents=True, exist_ok=True)
    generated = []
    for folder, mapping in TARGETS.items():
        for idx, (name, kind) in enumerate(mapping.items()):
            p = data_dir / folder / name
            p.parent.mkdir(parents=True, exist_ok=True)
            WRITERS[kind](p, SEED + idx + len(folder))
            generated.append(p)
    for idx, (name, kind) in enumerate(DECOYS.items()):
        p = data_dir / "A_COMMAND_ARCHIVE" / name
        p.parent.mkdir(parents=True, exist_ok=True)
        WRITERS[kind](p, SEED + 100 + idx)
        generated.append(p)
    # Stable, staggered modification times for repeatable lab trials.
    base_ts = 1_790_000_000
    for i, p in enumerate(generated):
        os.utime(p, (base_ts + i * 7, base_ts + i * 7))
    hashes = {str(p.relative_to(data_dir)): sha256_file(p) for p in sorted(generated)}
    results_root = data_dir.parent if data_dir.name == "MILITARY_DATA" else root
    results_dir = results_root / "results"
    results_dir.mkdir(parents=True, exist_ok=True)
    (results_dir / "baseline_hashes.json").write_text(json.dumps({
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "data_dir": str(data_dir), "sha256": hashes
    }, indent=2), encoding="utf-8")
    print(f"Generated {len(generated)} fictional files under {data_dir}")
    print(f"Hashes saved to {results_dir / 'baseline_hashes.json'}")
    return hashes

if __name__ == "__main__":
    base = getattr(config, "TEST_DATA_DIR", str(Path.home() / "MILITARY_DATA")) if config else str(Path.home())
    # If config points directly at MILITARY_DATA, generator recognizes it.
    generate_dataset(base)
